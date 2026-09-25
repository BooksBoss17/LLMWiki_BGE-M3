#!/usr/bin/env python
"""
MD → DOCX 转换脚本模板
功能：Markdown 转 Word，含 OMML 公式、双栏 A4 排版、PDF 导出
字体：宋体 + Times New Roman（用户卷子标准）

用法：
    python md2docx_template.py <input.md> [output.docx] [--columns 1|2] [--pages N]

如不指定 output，默认与 input 同名 .docx
"""

import re, os, sys, time, shutil, tempfile
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from lxml import etree

# ─── Import OMML converter ───
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'references'))
try:
    from latex_to_omml import LatexToOMML
except ImportError:
    raise ImportError("需要 latex_to_omml.py，请从 skill 的 references/ 目录复制")

# ═══════════════════════════════════════════════════════════════
# 配置参数
# ═══════════════════════════════════════════════════════════════
CONFIG = {
    'page_width_mm': 210,      # A4
    'page_height_mm': 297,
    'margin_top_cm': 0.8,
    'margin_bottom_cm': 0.8,
    'margin_left_cm': 1.0,
    'margin_right_cm': 1.0,
    # ⚠️ 用户卷子标准：宋体 + Times New Roman，不要用微软雅黑
    'font_cn': '宋体',          # 中文字体
    'font_en': 'Times New Roman', # 西文字体
    'font_size_pt': 7,         # 正文字号（讲义压缩；标准试卷用10.5）
    'line_spacing_pt': 9,
    'heading_sizes': {1: 11, 2: 8.5, 3: 7.5, 4: 7},
    'table_font_size': 6.5,
    'math_font_size': 8,
    'columns': 2,              # 双栏
    'target_pages': 2,         # 目标页数
    'column_spacing_pt': 6,    # 栏间距
}

# ═══════════════════════════════════════════════════════════════
# Helper functions
# ═══════════════════════════════════════════════════════════════
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def set_run_font(run, name=None, size=None):
    """Set run font: 宋体 for CJK, Times New Roman for Latin."""
    en_font = name or CONFIG['font_en']
    cn_font = CONFIG['font_cn']
    if size:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn('w:ascii'), en_font)
    rf.set(qn('w:hAnsi'), en_font)
    rf.set(qn('w:eastAsia'), cn_font)

def process_inline(text, run, converter):
    """Process **bold**, *italic*, `code`, $latex$ in text, apply to run."""
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$[^$]+\$)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run.bold = True
            run.add_text(part[2:-2])
            run.bold = False
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run.italic = True
            run.add_text(part[1:-1])
            run.italic = False
        elif part.startswith('`') and part.endswith('`'):
            run.font.name = 'Consolas'
            run.add_text(part[1:-1])
            run.font.name = CONFIG['font_en']
        elif part.startswith('$') and part.endswith('$'):
            omath = converter.convert_inline(part[1:-1])
            run._p.append(omath)
        else:
            run.add_text(part)

def add_heading(doc, text, level, converter):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2 if level <= 2 else 1)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = Pt(10)
    run = p.add_run()
    sz = CONFIG['heading_sizes'].get(level, 7)
    set_run_font(run, None, sz)
    run.bold = True
    # 黑色加粗（不用蓝色，与试卷一致）
    run.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if level == 2:
        pborder = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="CCCCCC"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pborder)
    run.add_text(text)

def add_para_with_math(doc, text, converter, size=None, bold=False, italic=False, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = Pt(CONFIG['line_spacing_pt'])
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    segments = re.split(r'(\$[^$]+\$)', text)
    for seg in segments:
        if not seg:
            continue
        if seg.startswith('$') and seg.endswith('$'):
            omath = converter.convert_inline(seg[1:-1])
            p._p.append(omath)
        else:
            sub_segs = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', seg)
            for ss in sub_segs:
                if not ss:
                    continue
                if ss.startswith('**') and ss.endswith('**'):
                    r = p.add_run(); set_run_font(r, None, size or CONFIG['font_size_pt'])
                    r.bold = True; r.add_text(ss[2:-2])
                elif ss.startswith('*') and ss.endswith('*') and not ss.startswith('**'):
                    r = p.add_run(); set_run_font(r, None, size or CONFIG['font_size_pt'])
                    r.italic = True; r.add_text(ss[1:-1])
                elif ss.startswith('`') and ss.endswith('`'):
                    r = p.add_run(); set_run_font(r, 'Consolas', (size or CONFIG['font_size_pt']) - 1)
                    r.add_text(ss[1:-1])
                else:
                    r = p.add_run(); set_run_font(r, None, size or CONFIG['font_size_pt'])
                    if bold: r.bold = True
                    if italic: r.italic = True
                    r.add_text(ss)
    return p

def add_display_math(doc, latex, converter):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = Pt(10)
    omathpara = converter.convert_display(latex)
    p._p.append(omathpara)

def add_table_from_md(doc, rows_md, converter):
    parsed = []
    for row in rows_md:
        if re.match(r'^[\s\|:-]+$', row):
            continue
        cells = [c.strip() for c in row.strip('|').split('|')]
        parsed.append(cells)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    mar = parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:left w:w="20" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="20" w:type="dxa"/></w:tblCellMar>')
    table._tbl.tblPr.append(mar)
    for ri, row_data in enumerate(parsed):
        for ci in range(ncols):
            cell_text = row_data[ci] if ci < len(row_data) else ''
            cell = table.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(8)
            segments = re.split(r'(\$[^$]+\$)', cell_text)
            for seg in segments:
                if not seg:
                    continue
                if seg.startswith('$') and seg.endswith('$'):
                    omath = converter.convert_inline(seg[1:-1])
                    p._p.append(omath)
                else:
                    sub = re.split(r'(\*\*[^*]+\*\*)', seg)
                    for ss in sub:
                        if not ss:
                            continue
                        if ss.startswith('**') and ss.endswith('**'):
                            r = p.add_run(); set_run_font(r, None, CONFIG['table_font_size'])
                            r.bold = True; r.add_text(ss[2:-2])
                        else:
                            r = p.add_run(); set_run_font(r, None, CONFIG['table_font_size'])
                            if ri == 0: r.bold = True
                            r.add_text(ss)
            if ri == 0:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E2F3" w:val="clear"/>')
                cell._tc.get_or_add_tcPr().append(shading)

def add_quote(doc, text, converter):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.line_spacing = Pt(8.5)
    pborder = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="8" w:space="3" w:color="666666"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pborder)
    segments = re.split(r'(\$[^$]+\$)', text)
    for seg in segments:
        if not seg:
            continue
        if seg.startswith('$') and seg.endswith('$'):
            omath = converter.convert_inline(seg[1:-1])
            p._p.append(omath)
        else:
            r = p.add_run(); set_run_font(r, None, CONFIG['table_font_size'])
            r.italic = True; r.add_text(seg)

def add_list_item(doc, text, converter, num=None, level=0, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(CONFIG['line_spacing_pt'])
    p.paragraph_format.left_indent = Cm(0.4 + level * 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    r = p.add_run(); set_run_font(r, None, CONFIG['font_size_pt'])
    r.bold = bold
    r.add_text(f'{num}. ' if num else '• ')
    segments = re.split(r'(\$[^$]+\$)', text)
    for seg in segments:
        if not seg:
            continue
        if seg.startswith('$') and seg.endswith('$'):
            omath = converter.convert_inline(seg[1:-1])
            p._p.append(omath)
        else:
            sub = re.split(r'(\*\*[^*]+\*\*)', seg)
            for ss in sub:
                if not ss:
                    continue
                if ss.startswith('**') and ss.endswith('**'):
                    r = p.add_run(); set_run_font(r, None, CONFIG['font_size_pt'])
                    r.bold = True; r.add_text(ss[2:-2])
                else:
                    r = p.add_run(); set_run_font(r, None, CONFIG['font_size_pt'])
                    r.bold = bold; r.add_text(ss)

# ═══════════════════════════════════════════════════════════════
# Main conversion function
# ═══════════════════════════════════════════════════════════════
def convert_md_to_docx(md_path, docx_path=None, columns=None, target_pages=None,
                       skip_sections=None):
    """
    Convert Markdown to DOCX with OMML formulas.
    
    Args:
        md_path: Input .md file path
        docx_path: Output .docx path (default: same as input with .docx)
        columns: 1 or 2 (default: from CONFIG)
        target_pages: Target page count (default: from CONFIG)
        skip_sections: Set of section titles to skip (e.g. {'考情分析', '来源标注'})
    """
    columns = columns or CONFIG['columns']
    target_pages = target_pages or CONFIG['target_pages']
    skip_sections = skip_sections or set()
    
    if docx_path is None:
        docx_path = md_path.rsplit('.', 1)[0] + '.docx'
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md = f.read()
    
    converter = LatexToOMML()
    
    # Create document
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(CONFIG['page_width_mm'])
    section.page_height = Mm(CONFIG['page_height_mm'])
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(CONFIG['margin_top_cm'])
    section.bottom_margin = Cm(CONFIG['margin_bottom_cm'])
    section.left_margin = Cm(CONFIG['margin_left_cm'])
    section.right_margin = Cm(CONFIG['margin_right_cm'])
    # ⚠️ DO NOT set columns in python-docx — use Word COM in post-process
    
    # Normal style: 宋体 + TNR
    style = doc.styles['Normal']
    style.font.name = CONFIG['font_en']
    style.font.size = Pt(CONFIG['font_size_pt'])
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0.5)
    style.paragraph_format.line_spacing = Pt(CONFIG['line_spacing_pt'])
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), CONFIG['font_en'])
    rfonts.set(qn('w:hAnsi'), CONFIG['font_en'])
    rfonts.set(qn('w:eastAsia'), CONFIG['font_cn'])
    
    # Parse MD
    lines = md.split('\n')
    i = 0
    in_code = False
    code_buf = []
    in_tbl = False
    tbl_rows = []
    skip_mode = False
    
    while i < len(lines):
        line = lines[i]
        
        # Code blocks
        if line.strip().startswith('```'):
            if in_code:
                content = '\n'.join(code_buf)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.line_spacing = Pt(8)
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5" w:val="clear"/>')
                p._p.get_or_add_pPr().append(shading)
                r = p.add_run(); set_run_font(r, 'Consolas', 6)
                for cl in code_buf:
                    if cl.strip():
                        r.add_text(cl + '\n')
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        
        # Tables
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            if not in_tbl:
                in_tbl = True
                tbl_rows = []
            tbl_rows.append(line.strip())
            i += 1
            continue
        else:
            if in_tbl:
                add_table_from_md(doc, tbl_rows, converter)
                tbl_rows = []
                in_tbl = False
        
        # Skip sections
        if line.startswith('## ') and any(s in line for s in skip_sections):
            skip_mode = True
            i += 1
            continue
        if skip_mode and line.startswith('## '):
            skip_mode = False
        if skip_mode:
            i += 1
            continue
        
        # HTML tags
        if line.strip().startswith('<details>') or line.strip() == '</details>':
            i += 1
            continue
        if line.strip().startswith('<summary>'):
            txt = re.sub(r'</?summary>', '', line.strip())
            add_para_with_math(doc, f'【{txt}】', converter, size=6.5, bold=True)
            i += 1
            continue
        
        # Headings
        if line.startswith('#### '):
            add_heading(doc, line[5:].strip(), 4, converter)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 3, converter)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), 2, converter)
        elif line.startswith('# '):
            add_heading(doc, line[2:].strip(), 1, converter)
        # Display math
        elif line.strip().startswith('$$'):
            mc = line.strip()[2:]
            if mc.endswith('$$'):
                add_display_math(doc, mc[:-2], converter)
            else:
                ml = [mc]
                i += 1
                while i < len(lines) and not lines[i].strip().endswith('$$'):
                    ml.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    ml.append(lines[i].strip()[:-2])
                add_display_math(doc, ' '.join(m for m in ml if m), converter)
        # Blockquote
        elif line.startswith('> '):
            add_quote(doc, line[2:].strip(), converter)
        # List items
        elif re.match(r'^[\s]*[-*]\s', line):
            txt = re.sub(r'^[\s]*[-*]\s+', '', line)
            add_list_item(doc, txt, converter)
        elif re.match(r'^[\s]*\d+\.\s', line):
            txt = re.sub(r'^[\s]*\d+\.\s+', '', line)
            num = re.match(r'^[\s]*(\d+)\.', line).group(1)
            add_list_item(doc, txt, converter, num=num)
        # HR (skip)
        elif line.strip() in ('---', '***'):
            pass
        elif not line.strip():
            pass
        else:
            add_para_with_math(doc, line.strip(), converter)
        i += 1
    
    if tbl_rows:
        add_table_from_md(doc, tbl_rows, converter)
    
    # Remove trailing empty paragraphs
    while len(doc.paragraphs) > 0:
        last_p = doc.paragraphs[-1]
        if last_p.text.strip() == '' and len(last_p._p) <= 2:
            last_p._element.getparent().remove(last_p._element)
        else:
            break
    
    # Save
    doc.save(docx_path)
    sz = os.path.getsize(docx_path)
    print(f"DOCX saved: {docx_path} ({sz/1024:.1f} KB)")
    
    # Word COM post-process: set columns + export PDF
    post_process_with_word(docx_path, columns, target_pages)
    
    return docx_path


def post_process_with_word(docx_path, columns, target_pages):
    """Use Word COM to set columns and export PDF."""
    print(f"Word COM post-processing (columns={columns}, target={target_pages}p)...")
    import win32com.client
    
    pdf_path = docx_path.replace('.docx', '.pdf')
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = False
    
    try:
        doc = word.Documents.Open(docx_path)
        time.sleep(1)
        
        # Set columns
        if columns > 1:
            doc.Sections(1).PageSetup.TextColumns.SetCount(columns)
            doc.Sections(1).PageSetup.TextColumns.Spacing = CONFIG['column_spacing_pt']
        
        doc.Repaginate()
        pages = doc.ComputeStatistics(2)
        print(f"  Page count: {pages}")
        
        # Trim if over target
        if pages > target_pages:
            for _ in range(5):
                sel = word.Selection
                sel.EndKey(6)  # wdStory
                sel.TypeBackspace()
            doc.Repaginate()
            pages = doc.ComputeStatistics(2)
            print(f"  After trim: {pages}")
        
        doc.Save()
        
        # Export PDF (handle file lock)
        time.sleep(0.5)
        try:
            doc.ExportAsFixedFormat(pdf_path, ExportFormat=17, OpenAfterExport=False)
            print(f"  PDF exported: {pdf_path}")
        except Exception as e:
            # Fallback: use _v2 suffix (old PDF may be locked by reader)
            pdf_path = docx_path.replace('.docx', '_v2.pdf')
            doc.ExportAsFixedFormat(pdf_path, ExportFormat=17, OpenAfterExport=False)
            print(f"  PDF exported (v2): {pdf_path}")
        
        doc.Close(False)
    finally:
        try:
            word.Quit()
        except:
            pass
        time.sleep(0.5)
    
    # Verify
    try:
        import fitz
        pdf_doc = fitz.open(pdf_path)
        print(f"  PDF pages: {len(pdf_doc)}")
        for i, page in enumerate(pdf_doc):
            text = page.get_text()
            lines = [l for l in text.split('\n') if l.strip()]
            print(f"    Page {i+1}: {len(lines)} lines, {len(text)} chars")
        pdf_doc.close()
    except ImportError:
        print("  (PyMuPDF not available for verification)")


# ═══════════════════════════════════════════════════════════════
# CLI entry point
# ═══════════════════════════════════════════════════════════════
if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python md2docx_template.py <input.md> [output.docx] [--columns N] [--pages N]")
        sys.exit(1)
    
    md_path = sys.argv[1]
    docx_path = None
    columns = None
    target_pages = None
    
    for j, arg in enumerate(sys.argv[2:], 2):
        if arg == '--columns' and j+1 < len(sys.argv):
            columns = int(sys.argv[j+1])
        elif arg == '--pages' and j+1 < len(sys.argv):
            target_pages = int(sys.argv[j+1])
        elif not arg.startswith('--'):
            docx_path = arg
    
    # Default skip sections for 讲义 (not exam papers)
    skip = {'考情分析', '来源标注', '讲义说明', '适用教材', '复习范围', '课型'}
    convert_md_to_docx(md_path, docx_path, columns, target_pages, skip_sections=skip)
