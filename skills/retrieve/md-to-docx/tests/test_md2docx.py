import io
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from lxml import etree
from PIL import Image


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "md2docx.py"


class Md2DocxImageTests(unittest.TestCase):
    def run_validation_conversion(self, root: Path, markdown: str):
        markdown_path = root / "input.md"
        markdown_path.write_text(markdown, encoding="utf-8")
        output_path = root / "output.docx"
        completed = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                str(markdown_path),
                str(output_path),
                "--layout",
                "validation",
                "--no-pdf",
                "--officecli-qa",
                "off",
            ],
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        return completed, output_path

    def test_validation_layout_embeds_original_pixels_at_controlled_extent(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / "diagram.png"
            Image.new("RGB", (200, 100), "white").save(
                image_path,
                dpi=(100, 100),
            )
            markdown_path = root / "input.md"
            markdown_path.write_text(
                "# Image check\n\n![complete diagram](diagram.png)\n",
                encoding="utf-8",
            )
            output_path = root / "output.docx"

            completed = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(markdown_path),
                    str(output_path),
                    "--layout",
                    "validation",
                    "--no-pdf",
                    "--officecli-qa",
                    "off",
                ],
                text=True,
                encoding="utf-8",
                errors="replace",
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            document = Document(output_path)
            self.assertEqual(len(document.inline_shapes), 1)
            shape = document.inline_shapes[0]
            self.assertAlmostEqual(shape.width.inches, 200 / 150, places=2)
            self.assertAlmostEqual(shape.width / shape.height, 2.0, places=2)
            image_paragraph = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
            )
            self.assertNotEqual(
                image_paragraph.paragraph_format.line_spacing_rule,
                WD_LINE_SPACING.EXACTLY,
            )
            self.assertIn("complete diagram", document.part.blob.decode("utf-8"))
            with zipfile.ZipFile(output_path) as archive:
                media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
                with Image.open(io.BytesIO(archive.read(media_name))) as embedded:
                    self.assertEqual(embedded.size, (200, 100))

    def test_validation_layout_emits_schema_ordered_radical_without_heading_border(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Formula check\n\n## Section\n\n$\\sqrt{x}$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            }
            radical = xml.xpath("//m:rad", namespaces=namespaces)[0]
            self.assertEqual(
                [child.tag for child in radical],
                [
                    f"{{{namespaces['m']}}}radPr",
                    f"{{{namespaces['m']}}}deg",
                    f"{{{namespaces['m']}}}e",
                ],
            )
            self.assertEqual(
                radical[0][0].tag,
                f"{{{namespaces['m']}}}degHide",
            )
            self.assertFalse(xml.xpath("//w:pPr/w:pBdr", namespaces=namespaces))

    def test_nested_fraction_preserves_outer_parse_state(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Formula check\n\n$\\frac{a+\\frac{b}{c}}{d}+x$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            }
            formulas = xml.xpath("//m:oMath", namespaces=namespaces)
            self.assertEqual(len(formulas), 1)
            self.assertEqual(len(formulas[0].xpath(".//m:f", namespaces=namespaces)), 2)
            top_level_tags = [etree.QName(child).localname for child in formulas[0]]
            self.assertEqual(top_level_tags, ["f", "r", "r"])
            self.assertEqual(
                "".join(formulas[0].xpath("./m:r/m:t/text()", namespaces=namespaces)),
                "+x",
            )

    def test_common_physics_operators_render_as_symbols_not_command_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Formula check\n\n$\\int+\\sum+\\varepsilon+\\sim+\\triangle+\\ll+\\cdots+\\arcsin x+{\\rm Pa}$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            }
            formula_text = "".join(xml.xpath("//m:oMath//m:t/text()", namespaces=namespaces))
            self.assertEqual(formula_text, "∫+∑+ε+∼+△+≪+⋯+arcsin x+ Pa")
            for command_name in ("int", "sum", "varepsilon", "triangle", "cdots", "rm"):
                self.assertNotIn(command_name, formula_text)

    def test_fraction_paragraphs_do_not_use_exact_line_spacing(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Formula spacing\n\nInline $\\frac{a}{b}$ text.\n\n$$\\frac{c}{d}$$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            document = Document(output_path)
            formula_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph._p.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}f')
            ]
            self.assertEqual(len(formula_paragraphs), 2)
            for paragraph in formula_paragraphs:
                self.assertNotEqual(
                    paragraph.paragraph_format.line_spacing_rule,
                    WD_LINE_SPACING.EXACTLY,
                )

    def test_math_style_group_preserves_native_superscript(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Units\n\n$10\\,\\mathrm{m/s^2}$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            }
            self.assertEqual(len(xml.xpath("//m:sSup", namespaces=namespaces)), 1)
            formula_text = "".join(xml.xpath("//m:oMath//m:t/text()", namespaces=namespaces))
            self.assertNotIn("^", formula_text)

    def test_inline_double_dollar_and_emphasized_math_do_not_leak_markers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Mixed markup\n\nValue *$k$* and $$\\frac{a}{b}$$ and *A*.\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            }
            self.assertEqual(len(xml.xpath("//m:oMath", namespaces=namespaces)), 2)
            visible_text = "".join(xml.xpath("//w:t/text()", namespaces=namespaces))
            self.assertNotIn("$", visible_text)
            self.assertNotIn("*", visible_text)

    def test_optional_latex_whitespace_before_groups_is_normalized(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Spaced groups\n\n$\\dfrac {1}{2}mv_0^ {2}+\\sqrt {gh}+I_ {1}^ {2}$\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml")
            xml = etree.fromstring(document_xml)
            namespaces = {
                "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            }
            formula = xml.xpath("//m:oMath", namespaces=namespaces)[0]
            self.assertEqual("".join(formula.xpath(".//m:f/m:num//m:t/text()", namespaces=namespaces)), "1")
            self.assertEqual("".join(formula.xpath(".//m:f/m:den//m:t/text()", namespaces=namespaces)), "2")
            self.assertEqual("".join(formula.xpath(".//m:rad/m:e//m:t/text()", namespaces=namespaces)), "gh")
            superscripts = formula.xpath(".//m:sSup/m:sup | .//m:sSubSup/m:sup", namespaces=namespaces)
            self.assertEqual(["".join(node.xpath(".//m:t/text()", namespaces=namespaces)) for node in superscripts], ["2", "2"])

    def test_missing_referenced_image_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            markdown_path = root / "input.md"
            markdown_path.write_text(
                "# Missing image\n\nBefore ![required](missing.png) after\n",
                encoding="utf-8",
            )
            output_path = root / "output.docx"

            completed = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    str(markdown_path),
                    str(output_path),
                    "--layout",
                    "validation",
                    "--no-pdf",
                    "--officecli-qa",
                    "off",
                ],
                text=True,
                encoding="utf-8",
                errors="replace",
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )

            self.assertNotEqual(completed.returncode, 0)
            self.assertIn("referenced image does not exist", completed.stderr)
            self.assertFalse(output_path.exists())

    def test_validation_layout_caps_tall_image_extent_without_resampling(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / "portrait.png"
            Image.new("RGB", (600, 1800), "white").save(image_path)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Tall image\n\n![portrait](portrait.png)\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            document = Document(output_path)
            shape = document.inline_shapes[0]
            self.assertLessEqual(shape.height.cm, 12.02)
            self.assertAlmostEqual(shape.height / shape.width, 3.0, places=2)
            with zipfile.ZipFile(output_path) as archive:
                media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
                with Image.open(io.BytesIO(archive.read(media_name))) as embedded:
                    self.assertEqual(embedded.size, (600, 1800))

    def test_validation_layout_caps_wide_image_extent_without_resampling(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / "wide.png"
            Image.new("RGB", (3000, 600), "white").save(image_path)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Wide image\n\n![wide](wide.png)\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            document = Document(output_path)
            shape = document.inline_shapes[0]
            self.assertLessEqual(shape.width.cm, 15.02)
            self.assertAlmostEqual(shape.width / shape.height, 5.0, places=2)
            with zipfile.ZipFile(output_path) as archive:
                media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
                with Image.open(io.BytesIO(archive.read(media_name))) as embedded:
                    self.assertEqual(embedded.size, (3000, 600))

    def test_validation_layout_keeps_headings_with_next_and_uses_compact_vertical_margins(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            completed, output_path = self.run_validation_conversion(
                root,
                "# Question\n\n## 答案\n\nB\n",
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            document = Document(output_path)
            section = document.sections[0]
            self.assertAlmostEqual(section.top_margin.cm, 1.6, places=1)
            self.assertAlmostEqual(section.bottom_margin.cm, 1.6, places=1)
            headings = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text in {"Question", "答案"}
            ]
            self.assertEqual(len(headings), 2)
            self.assertTrue(all(p.paragraph_format.keep_with_next for p in headings))


if __name__ == "__main__":
    unittest.main()
