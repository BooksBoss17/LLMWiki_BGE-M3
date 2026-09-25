#!/usr/bin/env python
from __future__ import annotations

import argparse
from collections import defaultdict
import json
from pathlib import Path
import re
import subprocess
import sys
import textwrap
from typing import Any

from student_data_common import (
    ROOT,
    StudentDataError,
    add_mode_args,
    apply_migrations,
    audit,
    connect_db,
    main_guard,
    print_json,
    resolve_args_db,
)

REPO_ROOT = ROOT.parent
OUTPUT_DIR = REPO_ROOT / "output"


METHOD_ERROR_LABELS = {
    "concept_confusion": "概念混淆",
    "model_selection_error": "模型选择错误",
    "condition_misread": "条件读取错误",
    "direction_sign_error": "方向/符号错误",
    "formula_condition_error": "公式条件错误",
    "calculation_error": "计算错误",
    "process_incomplete": "过程不完整",
    "expression_unit_error": "表达/单位错误",
    "careless_or_omission": "审题或遗漏",
    "unknown": "未分类",
    "concept": "概念相关错误",
}


def pct(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{value * 100:.1f}%"


def signed_pct(value: float | None) -> str:
    if value is None:
        return "-"
    sign = "+" if value >= 0 else ""
    return f"{sign}{value * 100:.1f}%"


def mastery_level(value: float | None) -> str:
    if value is None:
        return "证据不足"
    if value >= 0.85:
        return "优势"
    if value >= 0.60:
        return "基本掌握"
    return "待巩固"


def method_label(value: str | None) -> str:
    text = (value or "").strip()
    if not text:
        return "-"
    return METHOD_ERROR_LABELS.get(text, text)


def advice_for(level: str, method_errors: set[str]) -> str:
    labels = [method_label(item) for item in sorted(method_errors) if item]
    if level == "优势":
        base = "保持题型迁移训练，优先挑战同知识点的综合变式。"
    elif level == "基本掌握":
        base = "用错题复盘补齐条件识别、模型选择和规范表达。"
    elif level == "待巩固":
        base = "先回到概念、模型条件和基础题，再推进到中档题。"
    else:
        base = "补充同类题证据后再判断掌握情况。"
    if labels:
        return f"{base}重点排查：{'、'.join(labels)}。"
    return base


def safe_cell(value: Any) -> str:
    return str(value or "-").replace("|", "｜").replace("\n", " ")


def short_join(values: set[str], *, limit: int = 2) -> str:
    cleaned = [item.strip() for item in sorted(values) if item and item.strip()]
    if not cleaned:
        return "-"
    return "；".join(cleaned[:limit])


def item_refs(item: dict[str, Any]) -> str:
    return "、".join(item.get("items") or []) or "-"


def describe_gap(gap: float | None) -> str:
    if gap is None:
        return "暂无班级对照"
    if gap >= 0.08:
        return f"高于班均 {signed_pct(gap)}"
    if gap <= -0.08:
        return f"低于班均 {signed_pct(gap)}"
    return f"接近班均（{signed_pct(gap)}）"


def top_items(items: list[dict[str, Any]], *, limit: int = 3, show_gap: bool = False) -> str:
    selected = items[:limit]
    if not selected:
        return "-"
    parts: list[str] = []
    for item in selected:
        gap_text = f"，{describe_gap(item.get('gap'))}" if show_gap else ""
        parts.append(f"{item['knowledge_point']}（题 {item_refs(item)}，{pct(item.get('mastery'))}{gap_text}）")
    return "；".join(parts)


def method_error_summary(items: list[dict[str, Any]], *, limit: int = 3) -> str:
    counts: dict[str, int] = defaultdict(int)
    for item in items:
        for error in item.get("method_errors") or []:
            label = method_label(error)
            if label != "-":
                counts[label] += 1
    if not counts:
        return "-"
    ranked = sorted(counts.items(), key=lambda pair: (-pair[1], pair[0]))
    return "、".join(f"{label}（{count}处）" for label, count in ranked[:limit])


def summarize_exam_profile(summaries: list[dict[str, Any]]) -> dict[str, Any]:
    weighted_student = 0.0
    weighted_class = 0.0
    total_weight = 0.0
    for item in summaries:
        weight = float(item.get("max_score_sum") or len(item.get("items") or []) or 1)
        mastery = item.get("mastery")
        class_avg = item.get("class_avg")
        if mastery is not None:
            weighted_student += float(mastery) * weight
            total_weight += weight
        if class_avg is not None:
            weighted_class += float(class_avg) * weight
    overall = weighted_student / total_weight if total_weight else None
    class_overall = weighted_class / total_weight if total_weight else None
    gap = overall - class_overall if overall is not None and class_overall is not None else None
    strong = [item for item in summaries if item["level"] == "优势"]
    basic = [item for item in summaries if item["level"] == "基本掌握"]
    weak = [item for item in summaries if item["level"] == "待巩固"]
    if overall is None:
        headline = "本次考试证据不足，暂不做整体判断。"
    elif overall >= 0.75:
        headline = "本次考试整体表现较稳，说明核心模型已有较好基础；后续重点不是大面积补缺，而是把临界题型做稳。"
    elif overall >= 0.55:
        headline = "本次考试整体处在中等区间，已经有能拿分的知识点，但综合题、条件辨析和模型迁移仍在拉低表现。"
    elif overall >= 0.35:
        headline = "本次考试整体偏弱，亮点集中在少数熟悉题型；复习时需要先补概念和基本模型，再推进到综合计算。"
    else:
        headline = "本次考试整体压力较大，当前最重要的是重建基础概念、模型条件和审题流程，暂时不宜直接堆综合难题。"
    return {
        "overall": overall,
        "class_overall": class_overall,
        "gap": gap,
        "headline": headline,
        "strong": strong,
        "basic": basic,
        "weak": weak,
    }


def sort_for_priority(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(
        items,
        key=lambda item: (
            item.get("mastery") is None,
            item.get("mastery") if item.get("mastery") is not None else 1,
            item.get("gap") if item.get("gap") is not None else 0,
            -float(item.get("max_score_sum") or len(item.get("items") or []) or 1),
        ),
    )


def sort_for_strength(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(
        items,
        key=lambda item: (
            item.get("mastery") is None,
            -(item.get("mastery") or 0),
            -(float(item.get("max_score_sum") or len(item.get("items") or []) or 1)),
            item["knowledge_point"],
        ),
    )


def task_type_for(item: dict[str, Any]) -> str:
    text = "；".join(
        [str(value) for value in list(item.get("key_points") or []) + list(item.get("difficulty_points") or [])]
    )
    kp = str(item.get("knowledge_point") or "")
    if "图像" in text or "图像" in kp:
        return "图像识别与物理量关系判断"
    if "变压器" in kp or "输电" in text or "线损" in text:
        return "变压器与远距离输电的链式计算"
    if "气体" in kp or "状态方程" in text or "等温" in text:
        return "气体状态方程与图像/实验情境"
    if "电磁感应" in kp or "法拉第" in kp or "导体" in text or "金属棒" in text:
        return "电磁感应过程分析与综合计算"
    if "带电粒子" in kp or "圆心角" in text or "轨迹" in text:
        return "带电粒子轨迹几何与分段运动"
    if "磁场" in kp:
        return "磁场方向、轨迹与基本模型判断"
    if "分子" in kp:
        return "分子力图像与势能变化判断"
    return "同知识点基础题到中档题"


def priority_sentence(item: dict[str, Any], *, urgent: bool) -> str:
    items = item_refs(item)
    task_type = task_type_for(item)
    key_point = short_join(item["key_points"], limit=1)
    difficulty = short_join(item["difficulty_points"], limit=1)
    errors = method_error_summary([item], limit=2)
    gap_text = describe_gap(item.get("gap"))
    if urgent:
        return (
            f"- {item['knowledge_point']}（题 {items}，掌握率 {pct(item['mastery'])}，班均 {pct(item['class_avg'])}，{gap_text}）："
            f"这是当前最需要先补的方向。训练不要从综合难题开始，先用 2-3 道基础题把 `{task_type}` 的概念、条件和符号方向说清楚，"
            f"再做同模型中档题。对应本卷重点：{key_point}；难点：{difficulty}；错因侧重：{errors}。"
        )
    return (
        f"- {item['knowledge_point']}（题 {items}，掌握率 {pct(item['mastery'])}，{gap_text}）："
        f"这一类不是完全不会，更像是稳定性不足。建议围绕 `{task_type}` 做一组同模型变式，"
        f"每题复盘条件读取、公式适用条件和规范表达，目标是把会做的题稳定做全。"
    )


def output_path(value: str | None, *, assessment_id: str) -> Path:
    default_name = f"{assessment_id}_student_mastery.md"
    out = Path(value) if value else Path(default_name)
    if out.suffix.lower() != ".md":
        out = out.with_suffix(".md")
    if not out.is_absolute():
        if len(out.parts) == 1:
            out = OUTPUT_DIR / out
        else:
            out = REPO_ROOT / out
    output_dir = OUTPUT_DIR
    resolved = out.resolve()
    output_resolved = output_dir.resolve()
    if resolved != output_resolved and output_resolved not in resolved.parents:
        raise StudentDataError(f"output must be under {output_resolved}")
    out.parent.mkdir(parents=True, exist_ok=True)
    return out


def companion_paths(md_path: Path) -> dict[str, Path]:
    return {
        "md": md_path,
        "docx": md_path.with_suffix(".docx"),
        "pdf": md_path.with_suffix(".pdf"),
        "chart": md_path.with_name(f"{md_path.stem}_mastery_chart.png"),
        "officecli_qa": OUTPUT_DIR / f"qa_{md_path.stem}",
    }


def clean_inline_markdown(text: str) -> str:
    cleaned = re.sub(r"`([^`]*)`", r"\1", text)
    cleaned = cleaned.replace("**", "")
    return cleaned


def split_markdown_row(line: str) -> list[str]:
    return [clean_inline_markdown(cell.strip()) for cell in line.strip().strip("|").split("|")]


def short_chart_label(value: str, *, width: int = 14) -> str:
    parts = [part for part in str(value or "未标注知识点").split("/") if part]
    label = "/".join(parts[-2:]) if len(parts) >= 2 else (parts[0] if parts else "未标注知识点")
    return "\n".join(textwrap.wrap(label, width=width, break_long_words=False)) or label


def choose_student(
    conn,
    *,
    class_id: str,
    assessment_id: str,
    student_id: str | None,
    seat_no: str | None,
    pick: str,
) -> dict[str, Any]:
    if student_id and seat_no:
        raise StudentDataError("choose either --student-id or --seat-no, not both")
    clauses = ["a.class_id = ?", "a.assessment_id = ?"]
    params: list[Any] = [class_id, assessment_id]
    if student_id:
        clauses.append("sar.student_id = ?")
        params.append(student_id)
        order_by = "sar.student_id ASC"
    elif seat_no:
        clauses.append("cm.seat_no = ?")
        params.append(str(seat_no))
        order_by = "sar.student_id ASC"
    elif pick == "random":
        order_by = "RANDOM()"
    else:
        order_by = "sar.student_id ASC"

    row = conn.execute(
        f"""
        SELECT
            sar.student_id,
            s.pseudonym,
            cm.seat_no,
            c.class_id,
            c.class_name,
            a.assessment_id,
            a.title AS assessment_title,
            a.assessed_on
        FROM student_assessment_results sar
        JOIN assessments a ON a.assessment_id = sar.assessment_id
        JOIN classes c ON c.class_id = a.class_id
        JOIN students s ON s.student_id = sar.student_id
        LEFT JOIN class_memberships cm
            ON cm.student_id = sar.student_id
            AND cm.class_id = a.class_id
            AND (cm.term_id = a.term_id OR a.term_id IS NULL OR cm.term_id IS NULL)
        WHERE {" AND ".join(clauses)}
        ORDER BY {order_by}
        LIMIT 1
        """,
        params,
    ).fetchone()
    if not row:
        target = student_id or (f"{class_id}/seat:{seat_no}" if seat_no else f"{class_id}/{assessment_id}")
        raise StudentDataError(f"no safe assessment result found for {target}")
    return dict(row)


def fetch_item_rows(conn, *, assessment_id: str, student_id: str) -> list[dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT
            ai.item_no,
            ai.display_order,
            ai.max_score,
            ai.key_point,
            ai.difficulty_point,
            q.question_id,
            kp.kp_id,
            kp.title AS knowledge_point,
            kp.wiki_path,
            r.score,
            r.method_error_type,
            r.error_type,
            CASE
                WHEN ai.max_score > 0 AND r.score IS NOT NULL
                THEN r.score / ai.max_score
                ELSE NULL
            END AS score_rate,
            AVG(
                CASE
                    WHEN ai.max_score > 0 AND cr.score IS NOT NULL
                    THEN cr.score / ai.max_score
                    ELSE NULL
                END
            ) AS class_avg_score_rate
        FROM assessment_items ai
        LEFT JOIN questions q ON q.question_id = ai.question_id
        LEFT JOIN question_knowledge_points qkp ON qkp.question_id = q.question_id
        LEFT JOIN knowledge_points kp ON kp.kp_id = qkp.kp_id
        JOIN student_item_results r
            ON r.assessment_item_id = ai.assessment_item_id
            AND r.student_id = ?
        LEFT JOIN student_item_results cr
            ON cr.assessment_item_id = ai.assessment_item_id
        WHERE ai.assessment_id = ?
        GROUP BY
            ai.assessment_item_id, ai.item_no, ai.display_order, ai.max_score,
            ai.key_point, ai.difficulty_point, q.question_id,
            kp.kp_id, kp.title, kp.wiki_path, r.score,
            r.method_error_type, r.error_type
        ORDER BY ai.display_order, ai.item_no, kp.kp_id
        """,
        (student_id, assessment_id),
    ).fetchall()
    result = [dict(row) for row in rows]
    if not result:
        raise StudentDataError(f"no item evidence found for {student_id} in {assessment_id}")
    return result


def summarize_by_knowledge(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "kp_id": "",
            "knowledge_point": "未标注知识点",
            "wiki_path": None,
            "items": [],
            "student_score_sum": 0.0,
            "max_score_sum": 0.0,
            "class_rate_weighted_sum": 0.0,
            "class_rate_weight": 0.0,
            "method_errors": set(),
            "key_points": set(),
            "difficulty_points": set(),
        }
    )
    for row in rows:
        group_key = row.get("kp_id") or f"unmapped:{row.get('question_id') or row.get('item_no')}"
        group = groups[group_key]
        group["kp_id"] = row.get("kp_id") or ""
        group["knowledge_point"] = row.get("knowledge_point") or "未标注知识点"
        group["wiki_path"] = row.get("wiki_path")
        group["items"].append(row.get("item_no") or "")
        max_score = float(row.get("max_score") or 0)
        score = float(row.get("score") or 0)
        group["student_score_sum"] += score
        group["max_score_sum"] += max_score
        class_rate = row.get("class_avg_score_rate")
        if class_rate is not None and max_score > 0:
            group["class_rate_weighted_sum"] += float(class_rate) * max_score
            group["class_rate_weight"] += max_score
        method_error = row.get("method_error_type") or row.get("error_type")
        if method_error:
            group["method_errors"].add(str(method_error))
        if row.get("key_point"):
            group["key_points"].add(str(row["key_point"]))
        if row.get("difficulty_point"):
            group["difficulty_points"].add(str(row["difficulty_point"]))

    summaries: list[dict[str, Any]] = []
    for group in groups.values():
        max_sum = group["max_score_sum"]
        mastery = group["student_score_sum"] / max_sum if max_sum > 0 else None
        class_avg = (
            group["class_rate_weighted_sum"] / group["class_rate_weight"]
            if group["class_rate_weight"] > 0
            else None
        )
        level = mastery_level(mastery)
        summaries.append(
            {
                "kp_id": group["kp_id"],
                "knowledge_point": group["knowledge_point"],
                "wiki_path": group["wiki_path"],
                "items": [item for item in group["items"] if item],
                "student_score_sum": group["student_score_sum"],
                "max_score_sum": group["max_score_sum"],
                "mastery": mastery,
                "class_avg": class_avg,
                "gap": mastery - class_avg if mastery is not None and class_avg is not None else None,
                "level": level,
                "method_errors": group["method_errors"],
                "key_points": group["key_points"],
                "difficulty_points": group["difficulty_points"],
                "advice": advice_for(level, group["method_errors"]),
            }
        )
    return sorted(summaries, key=lambda item: (item["mastery"] is None, item["mastery"] or 0, item["knowledge_point"]))


def generate_mastery_chart(summaries: list[dict[str, Any]], chart_path: Path) -> None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise StudentDataError("mastery chart output requires matplotlib") from exc

    chart_path.parent.mkdir(parents=True, exist_ok=True)
    items = sorted(
        summaries,
        key=lambda item: (
            item.get("mastery") is None,
            item.get("mastery") if item.get("mastery") is not None else 1,
            item["knowledge_point"],
        ),
    )
    labels = [short_chart_label(item["knowledge_point"]) for item in items]
    mastery_values = [(item.get("mastery") or 0) * 100 for item in items]
    class_values = [None if item.get("class_avg") is None else item["class_avg"] * 100 for item in items]
    colors = []
    for item in items:
        if item["level"] == "优势":
            colors.append("#2e7d32")
        elif item["level"] == "基本掌握":
            colors.append("#f9a825")
        else:
            colors.append("#c62828")

    plt.rcParams["font.sans-serif"] = ["SimSun", "Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    height = max(5.2, 0.52 * len(items) + 1.6)
    fig, ax = plt.subplots(figsize=(10.5, height))
    positions = list(range(len(items)))
    ax.barh(positions, mastery_values, color=colors, alpha=0.86, label="该生掌握率")
    class_x = [value for value in class_values if value is not None]
    class_y = [idx for idx, value in enumerate(class_values) if value is not None]
    if class_x:
        ax.scatter(class_x, class_y, color="#455a64", s=32, marker="D", label="班级均值")
    for idx, value in enumerate(mastery_values):
        ax.text(min(value + 1.2, 98), idx, f"{value:.0f}%", va="center", fontsize=8)
    ax.set_yticks(positions)
    ax.set_yticklabels(labels, fontsize=8.5)
    ax.set_xlim(0, 100)
    ax.set_xlabel("掌握率（%）")
    ax.set_title("知识点掌握率对比", fontsize=13, fontweight="bold")
    ax.grid(axis="x", linestyle="--", alpha=0.28)
    ax.legend(loc="lower right")
    ax.invert_yaxis()
    fig.tight_layout()
    fig.savefig(chart_path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def set_run_font(run, *, size: float = 10.5, bold: bool = False) -> None:
    from docx.shared import Pt
    from docx.oxml.ns import qn

    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def add_docx_paragraph(doc, text: str, *, style: str | None = None, size: float = 10.5, bold: bool = False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(clean_inline_markdown(text))
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_markdown_table_to_docx(doc, rows: list[list[str]]) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    if len(rows) >= 2 and all(set(cell.replace(":", "").strip()) <= {"-"} for cell in rows[1]):
        rows = [rows[0], *rows[2:]]
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_cells = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        header_cells[idx].text = value
        header_cells[idx]._tc.get_or_add_tcPr().append(
            parse_xml(r'<w:shd {} w:val="clear" w:fill="D9EAF7"/>'.format(nsdecls("w")))
        )
    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row[: len(cells)]):
            cells[idx].text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=7.2)
                paragraph.paragraph_format.space_after = Pt(0)


def markdown_to_docx(markdown: str, *, md_path: Path, docx_path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Inches, Mm, Pt
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise StudentDataError("DOCX output requires python-docx") from exc

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    lines = markdown.splitlines()
    idx = 0
    while idx < len(lines):
        line = lines[idx].rstrip()
        if not line:
            idx += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].startswith("|"):
                table_lines.append(split_markdown_row(lines[idx]))
                idx += 1
            add_markdown_table_to_docx(doc, table_lines)
            continue
        image_match = re.match(r"!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)", line)
        if image_match:
            image_path = Path(image_match.group("path"))
            if not image_path.is_absolute():
                image_path = md_path.parent / image_path
            if image_path.exists():
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Inches(9.0))
            idx += 1
            continue
        if line.startswith("# "):
            paragraph = add_docx_paragraph(doc, line[2:], size=16, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            add_docx_paragraph(doc, line[3:], size=12.5, bold=True)
        elif line.startswith("### "):
            add_docx_paragraph(doc, line[4:], size=11, bold=True)
        elif line.startswith("- "):
            add_docx_paragraph(doc, line[2:], style="List Bullet", size=10.0)
        else:
            add_docx_paragraph(doc, line, size=10.0)
        idx += 1
    doc.save(docx_path)


def export_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    try:
        import win32com.client
    except ImportError as exc:
        raise StudentDataError("PDF output requires Word COM / pywin32 on Windows") from exc

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_pdf = pdf_path.with_name(f"{pdf_path.stem}.__tmp__.pdf")
    if tmp_pdf.exists():
        tmp_pdf.unlink()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.ExportAsFixedFormat(str(tmp_pdf.resolve()), ExportFormat=17, OpenAfterExport=False)
        doc.Close(False)
        doc = None
        if pdf_path.exists():
            pdf_path.unlink()
        tmp_pdf.replace(pdf_path)
    except Exception as exc:
        raise StudentDataError(f"PDF export failed: {exc}") from exc
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                try:
                    word.Application.Quit()
                except Exception:
                    pass


def inspect_pdf(pdf_path: Path) -> dict[str, Any]:
    try:
        import fitz
    except ImportError:
        return {}
    doc = fitz.open(str(pdf_path))
    try:
        blank_pages = 0
        for page in doc:
            if not page.get_text().strip():
                blank_pages += 1
        return {"pages": len(doc), "blank_pages": blank_pages}
    finally:
        doc.close()


def run_officecli_qa(policy: str, *, docx_path: Path, pdf_path: Path, out_dir: Path) -> dict[str, Any]:
    if policy == "off":
        return {"ok": True, "status": "skipped", "reason": "disabled"}
    qa_script = REPO_ROOT / "skills" / "_shared" / "scripts" / "officecli_doc_qa.py"
    out_dir.mkdir(parents=True, exist_ok=True)
    check = subprocess.run(
        [sys.executable, str(qa_script), "check", "--json"],
        cwd=str(REPO_ROOT),
        text=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if check.returncode != 0:
        payload = {
            "ok": policy != "required",
            "status": "skipped",
            "reason": "officecli_unavailable",
            "stdout": check.stdout,
            "stderr": check.stderr,
        }
        (out_dir / "officecli_qa.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        if policy == "required":
            raise StudentDataError("OfficeCLI QA required but OfficeCLI is unavailable")
        return payload

    qa = subprocess.run(
        [
            sys.executable,
            str(qa_script),
            "qa",
            "--docx",
            str(docx_path.resolve()),
            "--pdf",
            str(pdf_path.resolve()),
            "--out-dir",
            str(out_dir.resolve()),
            "--json",
        ],
        cwd=str(REPO_ROOT),
        text=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    try:
        payload = json.loads(qa.stdout)
    except json.JSONDecodeError:
        payload = {"ok": False, "status": "error", "stdout": qa.stdout, "stderr": qa.stderr}
    if qa.returncode != 0 and policy == "required":
        raise StudentDataError("OfficeCLI QA required but QA failed")
    return payload


def render_report(
    student: dict[str, Any],
    summaries: list[dict[str, Any]],
    rows: list[dict[str, Any]],
    *,
    chart_ref: str | None = None,
) -> str:
    profile = summarize_exam_profile(summaries)
    priority_items = sort_for_priority(profile["weak"])
    urgent_items = priority_items[:5]
    consolidation_items = (sort_for_priority(profile["basic"]) + priority_items[5:])[:5]
    strong_items = sort_for_strength(profile["strong"])
    best_available_items = sort_for_strength(profile["strong"] + profile["basic"])
    error_focus = method_error_summary(urgent_items or summaries, limit=3)

    lines = [
        "# 单生期中知识点掌握脱敏报告",
        "",
        "## 基本信息",
        "",
        f"- class_id: `{student['class_id']}`",
        f"- class_name: {student['class_name']}",
        f"- assessment_id: `{student['assessment_id']}`",
        f"- assessment_title: {student['assessment_title']}",
        f"- assessed_on: {student.get('assessed_on') or '-'}",
        f"- student_id: `{student['student_id']}`",
        f"- pseudonym: {student['pseudonym']}",
        f"- seat_no: {student.get('seat_no') or '-'}",
        "",
        "## 隐私边界",
        "",
        "本报告仅保留脱敏学生标识、座位号、知识点、小题得分率、班级均值得分率、方法错误类型和复习建议；不含个人身份字段、整卷统计字段、排序字段、原始证据材料或教师备注。",
        "",
        "## 本次考试总体判断",
        "",
        profile["headline"],
        "",
        f"- 小题加权掌握率：{pct(profile['overall'])}；班级同卷均值：{pct(profile['class_overall'])}；相对差距：{signed_pct(profile['gap'])}，整体属于{describe_gap(profile['gap'])}。",
        f"- 结构特点：优势知识点 {len(profile['strong'])} 个，基本掌握 {len(profile['basic'])} 个，待巩固 {len(profile['weak'])} 个。",
    ]

    if strong_items:
        lines.append(
            f"- 本次考得比较好的方向：{top_items(strong_items, limit=3)}。这些点可以作为后续综合题中的稳定得分基础。"
        )
    elif best_available_items:
        lines.append(
            f"- 本次还没有达到优势档的知识点，但相对较稳的是：{top_items(best_available_items, limit=3)}。复习时可以把这些题作为找回手感的入口。"
        )

    if urgent_items:
        lines.append(
            f"- 主要不足集中在：{top_items(urgent_items, limit=4, show_gap=True)}。这些点既影响本卷得分，也容易牵连后续综合题。"
        )
    if consolidation_items:
        lines.append(
            f"- 需要继续巩固的是：{top_items(consolidation_items, limit=4, show_gap=True)}。这些题型不是完全空白，但还没有稳定到可放心放过。"
        )
    if error_focus != "-":
        lines.append(f"- 错因侧重点：{error_focus}。复习时要把错因写进解题步骤，而不是只记结论。")

    lines.extend(
        [
            "",
            "## 复习建议",
            "",
            "这份建议按本次卷面的实际证据排序：先补待巩固且相对差距明显的知识点，再巩固基本掌握但不稳定的题型，最后用优势点做迁移训练。",
            "",
            "### 最需要复习",
            "",
        ]
    )
    if urgent_items:
        for item in urgent_items:
            lines.append(priority_sentence(item, urgent=True))
    else:
        lines.append("- 暂无特别薄弱项，建议直接进入巩固提升。")

    lines.extend(["", "### 需要巩固", ""])
    if consolidation_items:
        for item in consolidation_items:
            lines.append(priority_sentence(item, urgent=False))
    else:
        lines.append("- 当前没有明显处在临界区的知识点。")

    lines.extend(["", "### 保持优势", ""])
    if strong_items:
        for item in strong_items[:3]:
            lines.append(
                f"- {item['knowledge_point']}（题 {item_refs(item)}，掌握率 {pct(item['mastery'])}）：本次表现好，后续不宜只停留在会做基础题，建议做一两道综合情境题保持迁移能力。"
            )
    else:
        lines.append("- 本次没有达到优势档的知识点，先把最需要复习的内容补起来。")

    if chart_ref:
        lines.extend(
            [
                "",
                "## 知识点掌握可视化",
                "",
                "图中彩色柱为该生掌握率，灰色菱形点为班级均值得分率，便于快速判断哪些知识点低于班级参照。",
                "",
                f"![知识点掌握率柱状图]({chart_ref})",
            ]
        )

    lines.extend(
        [
            "",
            "## 知识点掌握清单",
            "",
            "| 知识点 | 题号 | 掌握率 | 班均得分率 | 差距 | 判断 | 重点建议 |",
            "| --- | --- | ---: | ---: | ---: | --- | --- |",
        ]
    )
    for item in summaries:
        lines.append(
            "| {kp} | {items} | {mastery} | {class_avg} | {gap} | {level} | {advice} |".format(
                kp=safe_cell(item["knowledge_point"]),
                items=safe_cell(item_refs(item)),
                mastery=pct(item["mastery"]),
                class_avg=pct(item["class_avg"]),
                gap=signed_pct(item["gap"]),
                level=safe_cell(item["level"]),
                advice=safe_cell(item["advice"]),
            )
        )

    lines.extend(
        [
            "",
            "## 逐题证据",
            "",
            "| 题号 | 知识点 | 得分率 | 班均得分率 | 重点 | 难点 | 方法错误类型 |",
            "| --- | --- | ---: | ---: | --- | --- | --- |",
        ]
    )
    for row in rows:
        error = row.get("method_error_type") or row.get("error_type")
        lines.append(
            "| {item_no} | {kp} | {rate} | {class_rate} | {key_point} | {difficulty_point} | {error} |".format(
                item_no=safe_cell(row.get("item_no") or "-"),
                kp=safe_cell(row.get("knowledge_point") or "未标注知识点"),
                rate=pct(row.get("score_rate")),
                class_rate=pct(row.get("class_avg_score_rate")),
                key_point=safe_cell(row.get("key_point") or "-"),
                difficulty_point=safe_cell(row.get("difficulty_point") or "-"),
                error=safe_cell(method_label(error)),
            )
        )
    return "\n".join(lines) + "\n"


def run() -> int:
    parser = argparse.ArgumentParser(description="Export a safe pseudonymized exam student mastery report.")
    add_mode_args(parser, default_dev=True)
    parser.add_argument("--class-id", required=True)
    parser.add_argument("--assessment-id", required=True)
    parser.add_argument("--student-id", help="pseudonymized student_id; omitted chooses one safe student")
    parser.add_argument("--seat-no", help="seat number within class; safe seat-only selection")
    parser.add_argument("--pick", choices=["first", "random"], default="first")
    parser.add_argument("--out", help="output Markdown path under repository output/; defaults to output/<assessment_id>_student_mastery.md")
    parser.add_argument(
        "--officecli-qa",
        choices=["auto", "off", "required"],
        default="auto",
        help="Run controlled OfficeCLI post-export QA: auto, off, or required.",
    )
    parser.add_argument("--actor", default="agent")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    path, real_data = resolve_args_db(args)
    out = output_path(args.out, assessment_id=args.assessment_id)
    paths = companion_paths(out)
    conn = connect_db(path, real_data=real_data)
    apply_migrations(conn)
    student = choose_student(
        conn,
        class_id=args.class_id,
        assessment_id=args.assessment_id,
        student_id=args.student_id,
        seat_no=args.seat_no,
        pick=args.pick,
    )
    rows = fetch_item_rows(conn, assessment_id=args.assessment_id, student_id=student["student_id"])
    summaries = summarize_by_knowledge(rows)
    generate_mastery_chart(summaries, paths["chart"])
    markdown = render_report(student, summaries, rows, chart_ref=paths["chart"].name)
    paths["md"].write_text(markdown, encoding="utf-8")
    markdown_to_docx(markdown, md_path=paths["md"], docx_path=paths["docx"])
    export_docx_to_pdf(paths["docx"], paths["pdf"])
    pdf_info = inspect_pdf(paths["pdf"])
    officecli_qa = run_officecli_qa(
        args.officecli_qa,
        docx_path=paths["docx"],
        pdf_path=paths["pdf"],
        out_dir=paths["officecli_qa"],
    )
    audit(
        conn,
        actor=args.actor,
        action="export",
        target=str(paths["md"].relative_to(REPO_ROOT)),
        purpose="safe exam student mastery export with chart/docx/pdf",
        row_count=len(rows),
        anonymized=True,
        real_data=real_data,
    )
    conn.close()

    payload = {
        "ok": True,
        "db": str(path),
        "out": str(paths["md"]),
        "outputs": {name: str(value) for name, value in paths.items()},
        "pdf": pdf_info,
        "officecli_qa": officecli_qa,
        "class_id": args.class_id,
        "assessment_id": args.assessment_id,
        "student_id": student["student_id"],
        "pseudonym": student["pseudonym"],
        "seat_no": student.get("seat_no"),
        "knowledge_points": len(summaries),
        "item_rows": len(rows),
        "real_data": real_data,
    }
    if args.json:
        print_json(payload)
    else:
        print(f"Wrote safe exam student mastery report: {paths['md']}")
        print(f"Wrote DOCX: {paths['docx']}")
        print(f"Wrote PDF: {paths['pdf']}")
        print(f"Wrote chart: {paths['chart']}")
        print(f"OfficeCLI QA: {officecli_qa.get('status')} ({paths['officecli_qa']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main_guard(run))
