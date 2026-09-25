#!/usr/bin/env python
"""Normalize a roster .xlsx into the safe StudentDataSQL roster CSV format.

The script reads full names only in process memory to derive family names. It
does not write full names, school numbers, exam numbers, or student-status
numbers to normalized outputs.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any

from student_data_common import IMPORTS_DIR, StudentDataError, assert_child_path, main_guard, print_json


DEFAULT_OUT_DIR = IMPORTS_DIR / "normalized"
FORBIDDEN_HEADERS = {
    "姓名",
    "学生姓名",
    "学号",
    "考号",
    "学籍号",
    "real_name",
    "school_number",
    "exam_number",
    "student_status_number",
}
COMPOUND_SURNAMES = (
    "欧阳",
    "太史",
    "端木",
    "上官",
    "司马",
    "东方",
    "独孤",
    "南宫",
    "万俟",
    "闻人",
    "夏侯",
    "诸葛",
    "尉迟",
    "公羊",
    "赫连",
    "澹台",
    "皇甫",
    "宗政",
    "濮阳",
    "公冶",
    "太叔",
    "申屠",
    "公孙",
    "慕容",
    "仲孙",
    "钟离",
    "长孙",
    "宇文",
    "司徒",
    "鲜于",
    "司空",
)

HEADER_ALIASES = {
    "seat_no": {"座号", "座位号", "班内学号", "班内学号（座位号）", "班内学号(座位号)", "seat_no"},
    "gender": {"性别", "gender"},
    "family_name": {"姓氏", "学生姓氏", "family_name"},
    "real_name": {"姓名", "学生姓名", "real_name", "name"},
    "class_id": {"class_id", "班级ID", "班级id"},
    "class_name": {"班级", "班级名称", "class_name"},
    "term_id": {"term_id", "学期ID", "学期id"},
}


def discover_xlsx(token: str) -> Path:
    desktop = Path.home() / "Desktop"
    candidates = [p for p in desktop.glob("*.xlsx") if token in p.name and not p.name.startswith("~$")]
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise StudentDataError(f"no .xlsx file containing token {token!r} found on Desktop")
    return candidates[0]


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text


def normalize_gender(value: str) -> str:
    text = (value or "").strip().upper()
    if text in {"M", "MALE", "男", "男生"}:
        return "M"
    if text in {"F", "FEMALE", "女", "女生"}:
        return "F"
    if text in {"U", "UNKNOWN", "未知"}:
        return "U"
    raise StudentDataError(f"unsupported gender value: {value!r}")


def seat_text(value: str) -> str:
    text = cell_text(value)
    if not text:
        raise StudentDataError("blank seat_no")
    return str(int(text)) if text.isdigit() else text


def family_name_from_real_name(real_name: str) -> str:
    name = re.sub(r"\s+", "", real_name or "")
    if not name:
        raise StudentDataError("blank real_name cannot derive family_name")
    for surname in COMPOUND_SURNAMES:
        if name.startswith(surname):
            return surname
    return name[0]


def header_map(row: tuple[Any, ...]) -> dict[str, int]:
    result: dict[str, int] = {}
    for idx, value in enumerate(row):
        text = cell_text(value)
        if not text:
            continue
        for canonical, aliases in HEADER_ALIASES.items():
            if text in aliases and canonical not in result:
                result[canonical] = idx
    return result


def find_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows[:20]):
        mapped = header_map(row)
        if "seat_no" in mapped and ("real_name" in mapped or "family_name" in mapped):
            return idx, mapped
    raise StudentDataError("could not find roster header row with seat_no and name/family_name columns")


def value_from(row: tuple[Any, ...], mapped: dict[str, int], key: str) -> str:
    idx = mapped.get(key)
    if idx is None or idx >= len(row):
        return ""
    return cell_text(row[idx])


def required_context(args: argparse.Namespace, mapped: dict[str, int]) -> None:
    missing = []
    if not args.term_id and "term_id" not in mapped:
        missing.append("term_id")
    if not args.class_id and "class_id" not in mapped:
        missing.append("class_id")
    if not args.class_name and "class_name" not in mapped:
        missing.append("class_name")
    if "gender" not in mapped:
        missing.append("gender")
    if missing:
        raise StudentDataError(
            "missing required roster fields: "
            + ", ".join(missing)
            + ". Stop and ask the user instead of guessing."
        )


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    fields = [
        "term_id",
        "school_year",
        "term_name",
        "class_id",
        "class_name",
        "grade_level",
        "student_id",
        "pseudonym",
        "family_name",
        "gender",
        "seat_no",
        "joined_on",
        "status",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run() -> int:
    parser = argparse.ArgumentParser(description="Preprocess roster .xlsx into safe StudentDataSQL roster CSV.")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--xlsx", help="source roster .xlsx path")
    source.add_argument("--desktop-token", help="discover latest Desktop .xlsx whose filename contains this token")
    parser.add_argument("--term-id", help="term id, for example 2025-2026-s2")
    parser.add_argument("--school-year", default="2025-2026")
    parser.add_argument("--term-name", default="second")
    parser.add_argument("--class-id", help="class id, for example g2c02")
    parser.add_argument("--class-name", help="class name, for example 高二02班")
    parser.add_argument("--grade-level", default="高二")
    parser.add_argument("--out", help="output CSV path under StudentDataSQL/runtime/imports")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise StudentDataError("openpyxl is required to read .xlsx roster files") from exc

    source_path = Path(args.xlsx).resolve() if args.xlsx else discover_xlsx(args.desktop_token).resolve()
    if not source_path.exists():
        raise StudentDataError(f"source xlsx does not exist: {source_path}")

    wb = load_workbook(source_path, read_only=True, data_only=True)
    ws = wb.active
    rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
    header_idx, mapped = find_header(rows)
    required_context(args, mapped)

    normalized: list[dict[str, str]] = []
    full_names_for_scan: list[str] = []
    seen_seats: set[str] = set()
    seen_ids: set[str] = set()

    for row in rows[header_idx + 1 :]:
        if not any(cell_text(value) for value in row):
            continue
        seat_no = seat_text(value_from(row, mapped, "seat_no"))
        if seat_no in seen_seats:
            raise StudentDataError(f"duplicate seat_no: {seat_no}")
        seen_seats.add(seat_no)

        gender = normalize_gender(value_from(row, mapped, "gender"))
        real_name = value_from(row, mapped, "real_name")
        family_name = value_from(row, mapped, "family_name") or family_name_from_real_name(real_name)
        if real_name:
            full_names_for_scan.append(real_name)

        term_id = args.term_id or value_from(row, mapped, "term_id")
        class_id = args.class_id or value_from(row, mapped, "class_id")
        class_name = args.class_name or value_from(row, mapped, "class_name")
        if not all([term_id, class_id, class_name, family_name, seat_no, gender]):
            raise StudentDataError("derived roster row has missing required fields; stop and ask the user")

        student_id = f"{term_id}_{class_id}_{gender}_seat{seat_no}"
        if student_id in seen_ids:
            raise StudentDataError(f"duplicate student_id: {student_id}")
        seen_ids.add(student_id)
        normalized.append(
            {
                "term_id": term_id,
                "school_year": args.school_year,
                "term_name": args.term_name,
                "class_id": class_id,
                "class_name": class_name,
                "grade_level": args.grade_level,
                "student_id": student_id,
                "pseudonym": f"{class_name}-{family_name}-{seat_no}号",
                "family_name": family_name,
                "gender": gender,
                "seat_no": seat_no,
                "joined_on": "",
                "status": "active",
            }
        )

    if not normalized:
        raise StudentDataError("no roster rows found")

    out = Path(args.out) if args.out else DEFAULT_OUT_DIR / f"{normalized[0]['class_id']}_{normalized[0]['term_id']}_roster.csv"
    if not out.is_absolute():
        out = ROOT.parent / out
    assert_child_path(out, IMPORTS_DIR)
    write_csv(out, normalized)

    output_text = out.read_text(encoding="utf-8-sig", errors="replace")
    leaked = [name for name in full_names_for_scan if len(name) > 1 and name in output_text]
    if leaked:
        out.unlink(missing_ok=True)
        raise StudentDataError("privacy check failed: normalized CSV contains full names")

    payload = {
        "ok": True,
        "source_name": source_path.name,
        "out": str(out),
        "row_count": len(normalized),
        "class_id": normalized[0]["class_id"],
        "term_id": normalized[0]["term_id"],
        "gender_counts": {g: sum(1 for row in normalized if row["gender"] == g) for g in ["F", "M", "U"]},
    }
    if args.json:
        print_json(payload)
    else:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main_guard(run))
